"""Document upload, parsing, chunking, and indexing services."""

from __future__ import annotations

import hashlib
import io
import re
import shutil
from contextlib import suppress
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import fitz
from fastapi import UploadFile
from pptx import Presentation
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.ai import AiProvider, IndexChunk, get_vector_index
from app.core.config import settings
from app.core.database import SessionLocal
from app.core.errors import AppError
from app.models import Document
from app.repositories import get_document_for_member, require_member, require_owner

ALLOWED_TYPES = {
    ".pdf": {"application/pdf"},
    ".pptx": {"application/vnd.openxmlformats-officedocument.presentationml.presentation"},
    ".md": {"text/markdown", "text/plain", "application/octet-stream"},
}


@dataclass(slots=True)
class ParsedSegment:
    text: str
    page: int | None = None
    slide: int | None = None
    section: str | None = None


@dataclass(slots=True)
class TextChunk:
    content: str
    page: int | None
    slide: int | None
    section: str | None


def list_documents(db: Session, course_id: str, user_id: str) -> list[Document]:
    require_member(db, course_id, user_id)
    return list(
        db.scalars(
            select(Document)
            .where(Document.course_id == course_id)
            .order_by(Document.created_at.desc())
        )
    )


def get_document(db: Session, document_id: str, user_id: str) -> Document:
    return get_document_for_member(db, document_id, user_id)


def create_document(
    db: Session,
    course_id: str,
    user_id: str,
    upload: UploadFile,
    content: bytes,
) -> Document:
    require_owner(db, course_id, user_id)
    suffix, original_name = validate_upload(
        upload.filename or "", upload.content_type or "", content
    )
    checksum = hashlib.sha256(content).hexdigest()
    existing = db.scalar(
        select(Document).where(Document.course_id == course_id, Document.checksum == checksum)
    )
    if existing:
        raise AppError(
            "DUPLICATE_DOCUMENT",
            "该文件已经上传，无需重复索引。",
            409,
            {"document_id": existing.id, "original_name": existing.original_name},
        )
    document = Document(
        course_id=course_id,
        original_name=original_name,
        storage_path="",
        mime_type=upload.content_type or _mime_for_suffix(suffix),
        size_bytes=len(content),
        checksum=checksum,
        status="pending",
        created_by=user_id,
    )
    db.add(document)
    db.flush()
    target_dir = settings.upload_dir / document.id
    target_dir.mkdir(parents=True, exist_ok=True)
    target = target_dir / f"source{suffix}"
    target.write_bytes(content)
    document.storage_path = str(target.resolve())
    db.commit()
    db.refresh(document)
    return document


def process_document(document_id: str) -> None:
    with SessionLocal() as db:
        document = db.get(Document, document_id)
        if not document:
            return
        document.status = "processing"
        document.error_message = None
        db.commit()
        try:
            segments = parse_document(Path(document.storage_path), document.original_name)
            if not segments:
                raise AppError(
                    "DOCUMENT_TEXT_EMPTY",
                    "未检测到可提取文本，扫描版 PDF 可能需要在后续版本接入 OCR。",
                    422,
                )
            chunks = chunk_segments(segments)
            if not chunks:
                raise AppError("DOCUMENT_TEXT_EMPTY", "文档没有可索引的有效文本。", 422)
            embeddings = AiProvider().embed([chunk.content for chunk in chunks])
            index_chunks = [
                IndexChunk(
                    id=_chunk_id(document.id, chunk, index),
                    course_id=document.course_id,
                    document_id=document.id,
                    file_name=document.original_name,
                    page=chunk.page,
                    slide=chunk.slide,
                    section=chunk.section,
                    content=chunk.content,
                    content_hash=hashlib.sha256(chunk.content.encode("utf-8")).hexdigest(),
                    embedding=embedding,
                )
                for index, (chunk, embedding) in enumerate(zip(chunks, embeddings, strict=True))
            ]
            vector_index = get_vector_index()
            vector_index.upsert_document(db, index_chunks)
            document.status = "ready"
            document.chunk_count = len(index_chunks)
            document.processed_at = datetime.utcnow()
            db.commit()
        except Exception as exc:  # noqa: BLE001
            db.rollback()
            document = db.get(Document, document_id)
            if document:
                document.status = "failed"
                document.error_message = _safe_error_message(exc)
                document.chunk_count = 0
                db.commit()
            with suppress(Exception):
                get_vector_index().delete_document(document_id)


def reindex_document(db: Session, document_id: str, user_id: str) -> Document:
    document = get_document_for_member(db, document_id, user_id)
    require_owner(db, document.course_id, user_id)
    document.status = "pending"
    document.error_message = None
    db.commit()
    return document


def delete_document(db: Session, document_id: str, user_id: str) -> None:
    document = get_document_for_member(db, document_id, user_id)
    require_owner(db, document.course_id, user_id)
    path = Path(document.storage_path)
    with suppress(Exception):
        get_vector_index().delete_document(document.id)
    db.delete(document)
    db.commit()
    if path.exists():
        shutil.rmtree(path.parent, ignore_errors=True)


def validate_upload(filename: str, mime_type: str, content: bytes) -> tuple[str, str]:
    if not filename or any(separator in filename for separator in ("/", "\\", "..")):
        raise AppError("UNSUPPORTED_FILE", "文件名不合法。", 415)
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED_TYPES:
        raise AppError("UNSUPPORTED_FILE", "仅支持 PDF、PPTX 和 Markdown 文件。", 415)
    if mime_type and mime_type not in ALLOWED_TYPES[suffix]:
        raise AppError("UNSUPPORTED_FILE", "文件 MIME 类型与扩展名不匹配。", 415)
    if len(content) > settings.max_upload_bytes:
        raise AppError("FILE_TOO_LARGE", f"文件不能超过 {settings.max_upload_mb} MB。", 413)
    if not content:
        raise AppError("VALIDATION_ERROR", "文件内容为空。", 422)
    if suffix == ".pdf" and not content.startswith(b"%PDF"):
        raise AppError("UNSUPPORTED_FILE", "文件头不是有效的 PDF。", 415)
    if suffix == ".pptx":
        if not content.startswith(b"PK"):
            raise AppError("UNSUPPORTED_FILE", "文件头不是有效的 PPTX。", 415)
        try:
            from zipfile import ZipFile

            with ZipFile(io.BytesIO(content)) as archive:
                names = set(archive.namelist())
                if "[Content_Types].xml" not in names or not any(
                    name.startswith("ppt/") for name in names
                ):
                    raise ValueError("not a pptx")
                uncompressed = sum(item.file_size for item in archive.infolist())
                if uncompressed > settings.max_upload_bytes * 8:
                    raise ValueError("archive too large")
        except (ValueError, OSError) as exc:
            raise AppError("UNSUPPORTED_FILE", "PPTX 文件结构无效。", 415) from exc
    return suffix, Path(filename).name


def parse_document(path: Path, original_name: str) -> list[ParsedSegment]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _parse_pdf(path)
    if suffix == ".pptx":
        return _parse_pptx(path)
    if suffix == ".md":
        return _parse_markdown(path)
    raise AppError("UNSUPPORTED_FILE", f"不支持的文件类型：{original_name}", 415)


def chunk_segments(
    segments: list[ParsedSegment],
    target_size: int = 700,
    max_size: int = 1000,
    overlap: int = 100,
) -> list[TextChunk]:
    chunks: list[TextChunk] = []
    for segment in segments:
        text = _clean_text(segment.text)
        if not text:
            continue
        start = 0
        while start < len(text):
            end = min(len(text), start + target_size)
            if end < len(text):
                boundary = _best_boundary(text, start + target_size // 2, end)
                if boundary > start:
                    end = boundary
            piece = text[start:end].strip()
            if piece:
                if len(piece) > max_size:
                    piece = piece[:max_size]
                    end = start + max_size
                chunks.append(
                    TextChunk(
                        content=piece,
                        page=segment.page,
                        slide=segment.slide,
                        section=segment.section,
                    )
                )
            if end >= len(text):
                break
            start = max(start + 1, end - overlap)
    return chunks


def _parse_pdf(path: Path) -> list[ParsedSegment]:
    segments: list[ParsedSegment] = []
    with fitz.open(path) as pdf:
        for page_number, page in enumerate(pdf, start=1):
            title = _first_line(page.get_text("text"))
            segments.append(
                ParsedSegment(text=page.get_text("text"), page=page_number, section=title)
            )
    return segments


def _parse_pptx(path: Path) -> list[ParsedSegment]:
    presentation = Presentation(str(path))
    segments: list[ParsedSegment] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        title = ""
        texts: list[str] = []
        for shape in slide.shapes:
            text = getattr(shape, "text", "")
            if text:
                if not title and getattr(shape, "has_text_frame", False):
                    title = text.strip().splitlines()[0]
                texts.append(text)
        segments.append(
            ParsedSegment(text="\n".join(texts), slide=slide_number, section=title or None)
        )
    return segments


def _parse_markdown(path: Path) -> list[ParsedSegment]:
    try:
        content = path.read_text(encoding="utf-8")
    except UnicodeDecodeError as exc:
        raise AppError("UNSUPPORTED_FILE", "Markdown 文件必须使用 UTF-8 编码。", 415) from exc
    if "\x00" in content:
        raise AppError("UNSUPPORTED_FILE", "Markdown 文件包含非法字符。", 415)
    segments: list[ParsedSegment] = []
    headings: list[str] = []
    paragraph: list[str] = []
    paragraph_number = 0
    code_block = False
    for raw_line in content.splitlines():
        line = raw_line.rstrip()
        if line.strip().startswith("```"):
            code_block = not code_block
            paragraph.append(line)
            continue
        if not code_block and re.match(r"^#{1,6}\s+", line):
            if paragraph:
                paragraph_number += 1
                segments.append(
                    ParsedSegment(
                        text="\n".join(paragraph),
                        section=" / ".join(headings) or f"段落 {paragraph_number}",
                    )
                )
                paragraph = []
            level = len(line) - len(line.lstrip("#"))
            headings = headings[: level - 1]
            headings.append(line.lstrip("#").strip())
        else:
            paragraph.append(line)
            if not line.strip() and paragraph:
                text = "\n".join(paragraph).strip()
                if text:
                    paragraph_number += 1
                    segments.append(
                        ParsedSegment(
                            text=text,
                            section=" / ".join(headings) or f"段落 {paragraph_number}",
                        )
                    )
                paragraph = []
    if paragraph:
        text = "\n".join(paragraph).strip()
        if text:
            segments.append(
                ParsedSegment(
                    text=text,
                    section=" / ".join(headings) or f"段落 {paragraph_number + 1}",
                )
            )
    return segments


def _clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _best_boundary(text: str, lower: int, upper: int) -> int:
    candidates = [
        text.rfind(mark, lower, upper) for mark in ("。", "；", "！", "？", "\n", ".", ";")
    ]
    boundary = max(candidates)
    return boundary + 1 if boundary >= 0 else upper


def _chunk_id(document_id: str, chunk: TextChunk, index: int) -> str:
    location = chunk.page or chunk.slide or 0
    content_hash = hashlib.sha256(chunk.content.encode("utf-8")).hexdigest()[:12]
    return f"{document_id}:{location}:{index}:{content_hash}"


def _first_line(text: str) -> str | None:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    return lines[0][:120] if lines else None


def _safe_error_message(exc: Exception) -> str:
    if isinstance(exc, AppError):
        return exc.message[:500]
    return "处理失败，请检查文件格式或稍后重试。"[:500]


def _mime_for_suffix(suffix: str) -> str:
    return {
        ".pdf": "application/pdf",
        ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        ".md": "text/markdown",
    }[suffix]
